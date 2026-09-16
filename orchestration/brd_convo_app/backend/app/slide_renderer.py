"""
slide_renderer.py — Render whole PowerPoint slides to PNG, so real flowcharts
and process flows make it into the generated SOW.

WHY THIS EXISTS
───────────────
image_extractor.py pulls *embedded* images out of a deck (the picture objects
and the raw zip media). That is the right tool for a screenshot or a pasted
diagram, but it structurally cannot capture the most valuable visuals in a
consulting deck, because those are not images at all: a PowerPoint flowchart is
normally drawn with native shapes — boxes, connectors, decision diamonds, swim
lanes, text frames — that exist only as OOXML geometry.

Measured on a real customer deck: the slide titled "Proposed - Hypercare support
process flow" carries 53 shapes, of which only 7 are pictures — and those 7 are
the little person/phone icons *inside* the flow. Extracting them yields seven
meaningless fragments and never the flowchart. The reviewer sees unrelated
pictures in the document and reasonably concludes the tool is inserting random
images.

The only way to get the actual diagram is to render the slide. On Windows with
Office installed we drive PowerPoint itself, which gives pixel-exact output —
the slide exactly as the author drew it.

SAFETY
──────
COM automation is run in a SEPARATE PROCESS with a timeout, never inline in the
web worker. PowerPoint can block indefinitely on a modal dialog (repair prompts,
font substitution, linked-media warnings); in-process that would wedge a worker
thread permanently. A subprocess can simply be killed.

Everything here degrades gracefully: no Windows, no Office, no pywin32, or a
render failure all return "no slides rendered" and leave the existing
embedded-image pipeline exactly as it was.
"""
from __future__ import annotations

import json
import logging
import os
import subprocess
import sys
import tempfile
from pathlib import Path
from typing import Dict, List, Optional

logger = logging.getLogger("SlideRenderer")

# Bound the work: rendering is ~1s/slide and each rendered slide costs a vision
# call afterwards. A deck's genuinely diagram-bearing slides are a small subset.
MAX_SLIDES_TO_RENDER = 24
RENDER_WIDTH_PX = 1600
# Whole-deck ceiling. Generous (a 95-slide deck rendered its picks in ~30s) but
# finite, so a hung or pathological deck can never stall an upload.
RENDER_TIMEOUT_S = 420

_DIAGRAM_WORDS = (
    "process flow", "workflow", "work flow", "to-be", "to be", "as-is", "as is",
    "architecture", "landscape", "timeline", "roadmap", "phases", "milestone",
    "governance", "operating model", "framework", "methodology", "approach",
    "lifecycle", "life cycle", "swimlane", "swim lane", "org chart",
    "organization", "structure", "integration", "data flow", "sequence",
    "escalation", "raci", "journey", "blueprint", "topology", "schedule",
)


def is_supported() -> bool:
    """True when whole-slide rendering can actually run on this machine."""
    if sys.platform != "win32":
        return False
    try:
        import win32com.client  # noqa: F401
    except Exception:
        return False
    return True


def pick_diagram_slides(pptx_path: str,
                        limit: int = MAX_SLIDES_TO_RENDER) -> List[int]:
    """Choose the slides worth rendering — 1-based slide numbers.

    A slide qualifies when it looks like a drawn diagram rather than a bullet
    list. The strongest single signal is the presence of CONNECTORS (<p:cxnSp>):
    an author only draws connectors to join boxes, which is exactly a flowchart
    or an architecture diagram. Shape density and title keywords catch the rest.
    """
    try:
        from pptx import Presentation
        from pptx.enum.shapes import MSO_SHAPE_TYPE
    except ImportError:
        logger.warning("python-pptx not installed — cannot pick diagram slides")
        return []

    try:
        prs = Presentation(pptx_path)
    except Exception as exc:
        logger.warning("Could not open %s for slide picking: %s", pptx_path, exc)
        return []

    def _walk(shapes):
        for s in shapes:
            try:
                if s.shape_type == MSO_SHAPE_TYPE.GROUP:
                    yield from _walk(s.shapes)
                    continue
            except Exception:
                pass
            yield s

    scored: List[tuple] = []
    for idx, slide in enumerate(prs.slides, start=1):
        try:
            shapes = list(_walk(slide.shapes))
        except Exception:
            continue

        n_conn = 0
        n_text = 0
        title_text = ""
        all_text = []
        for s in shapes:
            tag = getattr(getattr(s, "_element", None), "tag", "") or ""
            if tag.endswith("}cxnSp"):
                n_conn += 1
            try:
                if s.has_text_frame and s.text_frame.text.strip():
                    n_text += 1
                    all_text.append(s.text_frame.text.strip())
            except Exception:
                pass
        try:
            if slide.shapes.title is not None and slide.shapes.title.text:
                title_text = slide.shapes.title.text
        except Exception:
            pass

        blob = " ".join([title_text] + all_text).lower()
        kw = any(w in blob for w in _DIAGRAM_WORDS)
        n_shapes = len(shapes)

        # A pure bullet slide has many text frames but no connectors and few
        # non-text shapes; require real drawn structure, not just wordiness.
        n_nontext = n_shapes - n_text

        score = 0.0
        if n_conn >= 2:
            score += 5.0 + min(n_conn, 20) * 0.3
        if n_shapes >= 15:
            score += 2.0
        if n_nontext >= 8:
            score += 1.5
        if kw:
            score += 2.5
        # Title-only / near-empty slides are never diagrams.
        if n_shapes < 6:
            score = 0.0

        if score > 0:
            scored.append((score, idx))

    scored.sort(key=lambda t: (-t[0], t[1]))
    picked = sorted(i for _, i in scored[:limit])
    logger.info("Slide picker: %d slide(s) look like diagrams in %s -> %s",
                len(picked), Path(pptx_path).name, picked)
    return picked


# ─────────────────────────────────────────────────────────────────────────────
# Subprocess entry point — the only place COM is touched.
# ─────────────────────────────────────────────────────────────────────────────

def _render_worker(pptx_path: str, out_dir: str, slides: List[int],
                   width: int) -> Dict[str, str]:
    import pythoncom
    import win32com.client

    out = Path(out_dir)
    out.mkdir(parents=True, exist_ok=True)
    rendered: Dict[str, str] = {}

    pythoncom.CoInitialize()
    app = None
    pres = None
    try:
        app = win32com.client.Dispatch("PowerPoint.Application")
        pres = app.Presentations.Open(
            str(Path(pptx_path).resolve()), ReadOnly=True, WithWindow=False
        )
        total = pres.Slides.Count
        for n in slides:
            if n < 1 or n > total:
                continue
            dest = out / f"slide{n:04d}.png"
            try:
                pres.Slides(n).Export(str(dest), "PNG", width)
                if dest.exists() and dest.stat().st_size > 0:
                    rendered[str(n)] = str(dest)
            except Exception as exc:  # one bad slide must not kill the batch
                print(f"WARN slide {n}: {exc}", file=sys.stderr)
    finally:
        try:
            if pres is not None:
                pres.Close()
        except Exception:
            pass
        try:
            if app is not None:
                app.Quit()
        except Exception:
            pass
        pythoncom.CoUninitialize()
    return rendered


def render_slides(pptx_path: str, out_dir: str, slides: List[int],
                  width: int = RENDER_WIDTH_PX,
                  timeout: int = RENDER_TIMEOUT_S) -> Dict[int, str]:
    """Render the given 1-based slide numbers to PNG. Returns {slide_no: path}.

    Runs the COM work in a child process so a PowerPoint hang can be killed
    rather than wedging the caller. Returns {} on any failure.
    """
    if not slides:
        return {}
    if not is_supported():
        logger.info("Whole-slide rendering unavailable on this platform — "
                    "skipping (embedded-image extraction is unaffected).")
        return {}

    payload = json.dumps({
        "pptx": str(Path(pptx_path).resolve()),
        "out": str(Path(out_dir).resolve()),
        "slides": slides,
        "width": width,
    })

    try:
        proc = subprocess.run(
            [sys.executable, "-m", "app.slide_renderer", "--worker"],
            input=payload,
            capture_output=True,
            text=True,
            timeout=timeout,
            cwd=str(Path(__file__).resolve().parents[1]),
            env={**os.environ, "PYTHONIOENCODING": "utf-8"},
        )
    except subprocess.TimeoutExpired:
        logger.warning("Slide rendering timed out after %ss for %s — continuing "
                       "without whole-slide images.", timeout, Path(pptx_path).name)
        return {}
    except Exception as exc:
        logger.warning("Could not launch slide renderer for %s: %s",
                       Path(pptx_path).name, exc)
        return {}

    if proc.returncode != 0:
        logger.warning("Slide renderer failed (rc=%s) for %s: %s",
                       proc.returncode, Path(pptx_path).name,
                       (proc.stderr or "")[-400:])
        return {}

    try:
        raw = json.loads(proc.stdout.strip().splitlines()[-1])
    except Exception as exc:
        logger.warning("Slide renderer returned unparseable output for %s: %s",
                       Path(pptx_path).name, exc)
        return {}

    out = {int(k): v for k, v in raw.items()}
    logger.info("Rendered %d slide(s) from %s", len(out), Path(pptx_path).name)
    return out


def render_and_index_deck(pptx_path: str, output_dir: str, index_file: str,
                          source_doc: Optional[str] = None) -> Dict:
    """Render a deck's diagram slides and add them to the shared image index.

    Records are tagged ``render_kind="slide"`` so retrieval can prefer a whole
    rendered diagram over a fragment cropped out of the same slide.
    """
    from app.image_extractor import ImageExtractor

    src_name = source_doc or Path(pptx_path).name
    slides = pick_diagram_slides(pptx_path)
    if not slides:
        return {"rendered": 0, "indexed": 0}

    with tempfile.TemporaryDirectory(prefix="slide_render_") as tmp:
        paths = render_slides(pptx_path, tmp, slides)
        if not paths:
            return {"rendered": 0, "indexed": 0}

        extractor = ImageExtractor(output_dir=output_dir)
        records: List[Dict] = []
        for slide_no, png in sorted(paths.items()):
            try:
                raw = Path(png).read_bytes()
            except Exception:
                continue
            meta = extractor._save_image_record(
                raw, "png", source_doc=src_name, page_number=slide_no
            )
            if not meta:
                continue
            try:
                vision = extractor.describe_and_tag_image(
                    meta["file_path"],
                    source_context=f"{src_name}, slide {slide_no} (full slide)",
                    section_context="",
                )
            except Exception as exc:
                logger.warning("Vision failed for rendered slide %s: %s", slide_no, exc)
                vision = {}
            rec = {**meta, **vision}
            rec.pop("image_path", None)
            rec["render_kind"] = "slide"
            records.append(rec)

        added = extractor.append_to_index(records, index_file=index_file)

    try:
        from app.image_index import reload_image_index
        reload_image_index()
    except Exception:
        pass

    logger.info("Whole-slide rendering for %s: %d rendered, %d newly indexed",
                src_name, len(paths), len(added))
    return {"rendered": len(paths), "indexed": len(added)}


if __name__ == "__main__":
    if "--worker" in sys.argv:
        cfg = json.loads(sys.stdin.read())
        result = _render_worker(cfg["pptx"], cfg["out"], cfg["slides"], cfg["width"])
        print(json.dumps(result))
