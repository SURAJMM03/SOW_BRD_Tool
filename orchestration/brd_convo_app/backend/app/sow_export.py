"""
sow_export.py — render SOW Generate/Review output to a real .docx.

Uses its own Markdown → DOCX body renderer (_render_sow_body) rather than
chat_export._render_answer_docx. That helper was built for chat Q&A exports,
where the question occupies heading level 1, so it shifts every markdown
heading down one level and doesn't track bullet indentation depth. A SOW
draft's own headings (# for the report title, ## for section numbers, ###
for sub-sections) need to map 1:1 onto Word heading levels so the exported
outline/navigation pane matches the real 22-section structure, and its
nested bullets (e.g. under Section 22.1/22.2) need to stay nested rather
than flattening to one level, as the old shared renderer did.

Small parsing helpers (table detection, inline **bold**/`code` runs) are
reused from chat_export.py rather than duplicated.
"""

from __future__ import annotations

import re
from datetime import datetime
from pathlib import Path

from app.chat_export import EXPORT_DIR, _TBL_SEP, _add_runs, _row_cells

_HEAD = re.compile(r"^(#{1,4})\s+(.*)$")
_BULLET_INDENT = re.compile(r"^(\s*)[-*•]\s+(.*)$")
_NUMBERED = re.compile(r"^\s*\d+[.)]\s+(.*)$")

# Word's built-in nested-bullet styles — used up to 3 levels deep, which
# covers every level of nesting the SOW skill's own output actually produces.
_BULLET_STYLES = ["List Bullet", "List Bullet 2", "List Bullet 3"]


def _render_sow_body(doc, text: str) -> None:
    """Render one SOW draft/review's markdown into the docx document."""
    fences = (text or "").split("```")
    for fi, seg in enumerate(fences):
        if fi % 2 == 1:  # code fence body
            body = re.sub(r"^[a-zA-Z]*\n", "", seg)
            p = doc.add_paragraph()
            run = p.add_run(body.rstrip("\n"))
            run.font.name = "Consolas"
            continue

        lines = seg.replace("\r\n", "\n").split("\n")
        i = 0
        while i < len(lines):
            line = lines[i]

            # table
            if "|" in line and i + 1 < len(lines) and _TBL_SEP.match(lines[i + 1].strip()):
                header = _row_cells(line)
                i += 2
                rows = []
                while i < len(lines) and "|" in lines[i] and lines[i].strip():
                    rows.append(_row_cells(lines[i]))
                    i += 1
                table = doc.add_table(rows=1, cols=len(header))
                try:
                    table.style = "Table Grid"
                except Exception:
                    pass
                for ci, cell_text in enumerate(header):
                    cell_p = table.rows[0].cells[ci].paragraphs[0]
                    _add_runs(cell_p, cell_text)
                    for run in cell_p.runs:
                        run.bold = True
                for row in rows:
                    cells = table.add_row().cells
                    for ci in range(len(header)):
                        _add_runs(cells[ci].paragraphs[0], row[ci] if ci < len(row) else "")
                continue

            m = _HEAD.match(line.strip())
            if m:
                level = len(m.group(1))  # 1:1 with '#' count — no level shift
                doc.add_heading(re.sub(r"\*\*", "", m.group(2)), level=level)
                i += 1
                continue

            m = _BULLET_INDENT.match(line)
            if m:
                indent_ws, body_text = m.group(1), m.group(2)
                tabs = indent_ws.count("\t")
                spaces = len(indent_ws.replace("\t", ""))
                depth = min(tabs + spaces // 2, len(_BULLET_STYLES) - 1)
                p = doc.add_paragraph(style=_BULLET_STYLES[depth])
                _add_runs(p, body_text)
                i += 1
                continue

            m = _NUMBERED.match(line)
            if m:
                p = doc.add_paragraph(style="List Number")
                _add_runs(p, m.group(1))
                i += 1
                continue

            if line.strip():
                p = doc.add_paragraph()
                _add_runs(p, line.strip())
            i += 1


def sanitize_client_filename(client_name: str) -> str:
    """Turn a client name into a filesystem-safe token for `{Client}_SOW.docx`."""
    safe = re.sub(r"[^A-Za-z0-9 _-]", "", (client_name or "").strip())
    safe = re.sub(r"\s+", "_", safe).strip("_")
    return safe


def render_sow_docx(text: str, title: str, client_name: str = "") -> Path:
    from docx import Document

    EXPORT_DIR.mkdir(parents=True, exist_ok=True)
    safe_client = sanitize_client_filename(client_name)
    if safe_client:
        out = EXPORT_DIR / f"{safe_client}_SOW.docx"
    else:
        stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        out = EXPORT_DIR / f"SOW_{stamp}.docx"

    doc = Document()
    doc.add_heading(title, level=0)
    meta = doc.add_paragraph()
    meta.add_run(f"Generated {datetime.now().strftime('%Y-%m-%d %H:%M')}").italic = True

    _render_sow_body(doc, text)

    doc.save(str(out))
    return out
