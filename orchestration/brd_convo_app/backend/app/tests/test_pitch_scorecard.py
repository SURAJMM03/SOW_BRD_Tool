"""
Regression tests: Pitch Deck Scorecard maths
============================================
The scorecard exists to agree with a spreadsheet the Sales Enablement team
already uses. If the engine and the workbook ever disagree, the number on the
screen is wrong and nobody can tell by looking at it — so the sample scorecard
is pinned here as a golden fixture.

SAMPLE_RATINGS are the 57 ratings from "Scorecard sample.xlsx" (a review of a
102-slide Semtech proposal). The workbook computes a weighted total of
0.56825 and lands in the NOT READY band. So must we.

Run with:
    cd backend
    pytest app/tests/test_pitch_scorecard.py -v

pytest is not currently in requirements.txt, so this file also runs standalone:
    venv/Scripts/python.exe app/tests/test_pitch_scorecard.py
"""
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[2]))

from app import pitch_scorecard as ps  # noqa: E402

# Shorthand -> canonical rating, to keep the fixture readable.
_R = {
    "B": ps.RATING_BEST,
    "M": ps.RATING_MIN,
    "F": ps.RATING_FAIL,
    "N": ps.RATING_NA,
}

SAMPLE_RATINGS = {
    "1.1": "B", "1.2": "B", "1.3": "M", "1.4": "M",
    "2.1": "B", "2.2": "B", "2.3": "M",
    "3.1": "B", "3.2": "M", "3.3": "B", "3.4": "M",
    "4.1": "F", "4.2": "F", "4.3": "F", "4.4": "F",
    "5.1": "M", "5.2": "F", "5.3": "M", "5.4": "F", "5.5": "F",
    "6.1": "B", "6.2": "M", "6.3": "B", "6.4": "M", "6.5": "B", "6.6": "M",
    "7.1": "M", "7.2": "M", "7.3": "F", "7.4": "F", "7.5": "F",
    "8.1": "B", "8.2": "M", "8.3": "F", "8.4": "F",
    "9.1": "B", "9.2": "M", "9.3": "M", "9.4": "B",
    "10.1": "B", "10.2": "B", "10.3": "B", "10.4": "B",
    "11.1": "B", "11.2": "B", "11.3": "B", "11.4": "B",
    "12.1": "F", "12.2": "M", "12.3": "B", "12.4": "M",
    "13.1": "M", "13.2": "F", "13.3": "F",
    "14.1": "F", "14.2": "F", "14.3": "B",
}

# The workbook's own Summary!D4:D17.
SAMPLE_CATEGORY_AVERAGES = [
    1.5, 1.6666666666666667, 1.5, 0.0, 0.4, 1.5, 0.4, 0.75,
    1.5, 2.0, 2.0, 1.0, 0.3333333333333333, 0.6666666666666666,
]
SAMPLE_TOTAL = 0.5682500000000001

# Row spans on the "Pitch Checklist" sheet, from the source workbook's own
# B-column vertical merges.
EXPECTED_SPANS = [
    (4, 7), (8, 10), (11, 14), (15, 18), (19, 23), (24, 29), (30, 34),
    (35, 38), (39, 42), (43, 46), (47, 50), (51, 54), (55, 57), (58, 60),
]


def _answers(shorthand):
    return {cid: {"rating": _R[code]} for cid, code in shorthand.items()}


def _score(shorthand, weight_overrides=None):
    rubric = ps.load_rubric()
    store = ps._empty_store()
    if weight_overrides:
        store["weights"] = {str(k): v for k, v in weight_overrides.items()}
    weights = ps.effective_weights(rubric, store)
    return rubric, ps.compute_scores(rubric, _answers(shorthand), weights)


# ── Rubric integrity ─────────────────────────────────────────────────────────
class TestRubric:
    def test_shape(self):
        rubric = ps.load_rubric()
        assert len(rubric["categories"]) == 14
        assert len(rubric["criteria_index"]) == 57

    def test_weights_total_one(self):
        rubric = ps.load_rubric()
        assert round(sum(c["weight"] for c in rubric["categories"]), 6) == 1.0

    def test_row_spans_match_the_workbook(self):
        """The exported summary averages fixed cell ranges — these must not drift."""
        rubric = ps.load_rubric()
        spans = [(c["row_start"], c["row_end"]) for c in rubric["categories"]]
        assert spans == EXPECTED_SPANS

    def test_every_criterion_has_three_anchors(self):
        rubric = ps.load_rubric()
        for cid, crit in rubric["criteria_index"].items():
            for rating in (ps.RATING_FAIL, ps.RATING_MIN, ps.RATING_BEST):
                assert crit["anchors"].get(rating), f"{cid} missing {rating} anchor"

    def test_risks_anchors_were_corrected(self):
        """Category 14 shipped with pricing anchors copy-pasted in by mistake."""
        rubric = ps.load_rubric()
        for cid in ("14.1", "14.2", "14.3"):
            anchors = " ".join(rubric["criteria_index"][cid]["anchors"].values()).lower()
            assert "price not linked to value" not in anchors
            assert "loose linkage between cost" not in anchors

    def test_duplicate_criterion_is_category_scoped(self):
        """METHODOLOGY DEFINITION appears in two categories and must not collide."""
        rubric = ps.load_rubric()
        assert rubric["criteria_index"]["9.4"]["name"] == "METHODOLOGY DEFINITION"
        assert rubric["criteria_index"]["10.1"]["name"] == "METHODOLOGY DEFINITION"


# ── The golden fixture ───────────────────────────────────────────────────────
class TestSampleScorecard:
    def test_fixture_is_complete(self):
        assert len(SAMPLE_RATINGS) == 57

    def test_category_averages_match_the_workbook(self):
        _, result = _score(SAMPLE_RATINGS)
        got = [row["avg"] if row["avg"] is not None else 0.0
               for row in result["categories"]]
        for name, a, b in zip([r["name"] for r in result["categories"]],
                              got, SAMPLE_CATEGORY_AVERAGES):
            assert abs(a - b) < 1e-9, f"{name}: got {a}, workbook says {b}"

    def test_weighted_total_matches_the_workbook(self):
        _, result = _score(SAMPLE_RATINGS)
        assert abs(result["total"] - SAMPLE_TOTAL) < 1e-9

    def test_band_is_not_ready(self):
        _, result = _score(SAMPLE_RATINGS)
        assert result["band"] == ps.BAND_NOT_READY

    def test_fully_rated(self):
        _, result = _score(SAMPLE_RATINGS)
        assert result["unrated_count"] == 0
        assert result["complete"] is True


# ── Band ladder ──────────────────────────────────────────────────────────────
class TestBands:
    def test_ladder(self):
        assert ps.band_for(0.80) == ps.BAND_READY
        assert ps.band_for(0.95) == ps.BAND_READY
        assert ps.band_for(0.79) == ps.BAND_WORK
        assert ps.band_for(0.60) == ps.BAND_WORK
        assert ps.band_for(0.59) == ps.BAND_NOT_READY
        assert ps.band_for(0.01) == ps.BAND_NOT_READY
        assert ps.band_for(0.0) == ps.BAND_AWAITING


# ── Rating normalisation ─────────────────────────────────────────────────────
class TestNormalizeRating:
    def test_canonical(self):
        for rating in ps.VALID_RATINGS:
            assert ps.normalize_rating(rating) == rating

    def test_excel_emoji_forms(self):
        assert ps.normalize_rating("✅ Best Practice") == ps.RATING_BEST
        assert ps.normalize_rating("⚠️ Minimally Meets") == ps.RATING_MIN
        assert ps.normalize_rating("❌ Does Not Meet") == ps.RATING_FAIL

    def test_loose_model_output(self):
        assert ps.normalize_rating("best practice") == ps.RATING_BEST
        assert ps.normalize_rating("  Does Not Meet  ") == ps.RATING_FAIL
        assert ps.normalize_rating("N/A") == ps.RATING_NA
        assert ps.normalize_rating("na") == ps.RATING_NA

    def test_rejects_junk(self):
        for junk in ("", None, "Excellent", "2", "partially meets"):
            assert ps.normalize_rating(junk) is None


# ── N/A vs unrated — the one place we improve on the spreadsheet ─────────────
class TestNotApplicable:
    def test_na_criterion_leaves_the_average(self):
        """4 criteria, one N/A -> average over the other three."""
        ratings = dict(SAMPLE_RATINGS)
        ratings["1.1"] = "N"  # was B(2); 1.2=B(2) 1.3=M(1) 1.4=M(1)
        _, result = _score(ratings)
        structure = result["categories"][0]
        assert abs(structure["avg"] - (2 + 1 + 1) / 3) < 1e-9
        assert structure["na_count"] == 1

    def test_whole_category_na_is_renormalised_away(self):
        """PARTNERSHIPS N/A: its 10% redistributes, it does not score zero."""
        ratings = dict(SAMPLE_RATINGS)
        for cid in ("4.1", "4.2", "4.3", "4.4"):
            ratings[cid] = "N"
        _, result = _score(ratings)
        partnerships = result["categories"][3]
        assert partnerships["all_na"] is True
        assert partnerships["effective_weight"] == 0.0
        # Remaining weights renormalise to 1.00 ...
        assert abs(sum(r["effective_weight"] for r in result["categories"]) - 1.0) < 1e-9
        # ... so the total RISES: the sample scored PARTNERSHIPS 0 across the board.
        assert result["total"] > SAMPLE_TOTAL

    def test_unrated_is_not_treated_as_na(self):
        """An unscored category keeps its weight and drags the total down."""
        ratings = {k: v for k, v in SAMPLE_RATINGS.items()
                   if not k.startswith("11.")}  # AI & INNOVATION, all Best Practice
        _, result = _score(ratings)
        ai = result["categories"][10]
        assert ai["none_rated"] is True
        assert ai["unrated_count"] == 4
        assert ai["effective_weight"] > 0, "unrated must not renormalise away"
        assert result["complete"] is False
        assert result["total"] < SAMPLE_TOTAL

    def test_everything_na_does_not_divide_by_zero(self):
        ratings = {cid: "N" for cid in SAMPLE_RATINGS}
        _, result = _score(ratings)
        assert result["total"] == 0.0
        assert result["band"] == ps.BAND_AWAITING


# ── Reviewer weight overrides ────────────────────────────────────────────────
class TestWeightOverrides:
    def test_zeroing_a_category_redistributes_it(self):
        _, result = _score(SAMPLE_RATINGS, weight_overrides={4: 0.0})
        assert result["categories"][3]["effective_weight"] == 0.0
        assert abs(sum(r["effective_weight"] for r in result["categories"]) - 1.0) < 1e-9


# ── Workbook export ──────────────────────────────────────────────────────────
# app/tests -> app -> backend -> brd_convo_app -> orchestration -> <repo root>
SAMPLE_WORKBOOK = Path(__file__).resolve().parents[5] / "Scorecard sample.xlsx"


def _sample_summary():
    """A full summary dict for the sample ratings, via a throwaway store."""
    pid = "_test_pitch_export"
    store = ps._empty_store()
    store["answers"] = {cid: {"rating": _R[code]} for cid, code in SAMPLE_RATINGS.items()}
    store["deck_docs"] = ["Semtech_Proposal.pptx"]
    store["slide_counts"] = {"Semtech_Proposal.pptx": 102}
    ps.save_store(pid, store)
    try:
        return ps.summary(pid)
    finally:
        ps._store_path(pid).unlink(missing_ok=True)


class TestWorkbookExport:
    """The export must stay interchangeable with the team's own workbook.

    Three differences from "Scorecard sample.xlsx" are deliberate:
      * dims end at J60, not J61 — the sample's row 61 is entirely empty and
        exists only because a row height was set on it.
      * E34 and E41 lose a stray leading tab present in the source cells.
      * Category 14's anchors are the corrected ones (see PITCH_SKILL.md's
        changelog), so only categories 1-13 are compared verbatim.
    """

    def _built(self):
        from openpyxl import load_workbook
        from app import pitch_scorecard_xlsx as px
        return load_workbook(px.build_workbook(_sample_summary()), data_only=False)

    def _sample(self):
        from openpyxl import load_workbook
        return load_workbook(SAMPLE_WORKBOOK, data_only=False)

    def test_sheet_names(self):
        assert self._built().sheetnames == self._sample().sheetnames

    def test_merges_match(self):
        from app import pitch_scorecard_xlsx as px
        got, want = self._built(), self._sample()
        for sheet in (px.CHECKLIST_SHEET, px.SUMMARY_SHEET):
            assert (sorted(str(m) for m in got[sheet].merged_cells.ranges)
                    == sorted(str(m) for m in want[sheet].merged_cells.ranges)), sheet

    def test_headers_match(self):
        from app import pitch_scorecard_xlsx as px
        got, want = self._built(), self._sample()
        for col in "BCDEFGHIJ":
            assert (got[px.CHECKLIST_SHEET][f"{col}3"].value
                    == want[px.CHECKLIST_SHEET][f"{col}3"].value), f"checklist {col}3"
        for col in "BCDEFG":
            assert (got[px.SUMMARY_SHEET][f"{col}3"].value
                    == want[px.SUMMARY_SHEET][f"{col}3"].value), f"summary {col}3"

    def test_criteria_ratings_and_score_formulas_match(self):
        from app import pitch_scorecard_xlsx as px
        got, want = self._built()[px.CHECKLIST_SHEET], self._sample()[px.CHECKLIST_SHEET]
        for row in range(4, 61):
            assert got[f"D{row}"].value == want[f"D{row}"].value, f"criterion D{row}"
            assert got[f"H{row}"].value == want[f"H{row}"].value, f"rating H{row}"
            assert got[f"I{row}"].value == want[f"I{row}"].value, f"score formula I{row}"

    def test_anchors_match_for_uncorrected_categories(self):
        from app import pitch_scorecard_xlsx as px
        got, want = self._built()[px.CHECKLIST_SHEET], self._sample()[px.CHECKLIST_SHEET]
        for row in range(4, 58):  # categories 1-13; 14 is deliberately corrected
            for col in "EFG":
                a, b = got[f"{col}{row}"].value, want[f"{col}{row}"].value
                assert a == (b or "").strip(), f"anchor {col}{row}"

    def test_summary_names_and_weights_match(self):
        """With nothing marked N/A, effective weights equal the declared ones."""
        from app import pitch_scorecard_xlsx as px
        got, want = self._built()[px.SUMMARY_SHEET], self._sample()[px.SUMMARY_SHEET]
        for row in range(4, 18):
            assert got[f"B{row}"].value == want[f"B{row}"].value, f"B{row}"
            assert abs(got[f"C{row}"].value - want[f"C{row}"].value) < 1e-9, f"C{row}"

    def test_summary_formulas_average_the_same_ranges(self):
        """Formulas are IFERROR-wrapped (see test_all_na_category_...), so
        compare the ranges they act on rather than the literal strings."""
        from app import pitch_scorecard_xlsx as px
        got, want = self._built()[px.SUMMARY_SHEET], self._sample()[px.SUMMARY_SHEET]
        for row in range(4, 18):
            mine, theirs = got[f"D{row}"].value, want[f"D{row}"].value
            rng = theirs[theirs.index("!") + 1: theirs.index(")")]
            assert rng in mine, f"D{row} should average {rng}, got {mine}"
            assert got[f"E{row}"].value.count(f"D{row}/2") == 1, f"E{row}"
            assert got[f"F{row}"].value.count(f"E{row}*C{row}") == 1, f"F{row}"

    def test_total_and_band_cells_match(self):
        from app import pitch_scorecard_xlsx as px
        got, want = self._built()[px.SUMMARY_SHEET], self._sample()[px.SUMMARY_SHEET]
        for cell in ("B18", "C18", "F18", "G18", "B20"):
            assert got[cell].value == want[cell].value, cell

    def test_all_na_category_does_not_break_the_workbook(self):
        """Regression: AVERAGE over an all-N/A category is #DIV/0! in Excel.

        Column I holds the text "—" for every N/A criterion. Averaging a range
        that is entirely text is a division by zero, and unguarded that error
        propagates through E and F into SUM(F4:F17), destroying the total and
        the readiness band. Every summary formula must therefore be guarded.
        """
        from openpyxl import load_workbook
        from app import pitch_scorecard_xlsx as px
        ratings = dict(SAMPLE_RATINGS)
        for cid in ("4.1", "4.2", "4.3", "4.4"):
            ratings[cid] = "N"
        pid = "_test_pitch_na"
        store = ps._empty_store()
        store["answers"] = {c: {"rating": _R[v]} for c, v in ratings.items()}
        ps.save_store(pid, store)
        try:
            summary = ps.summary(pid)
        finally:
            ps._store_path(pid).unlink(missing_ok=True)

        ws = load_workbook(px.build_workbook(summary), data_only=False)[px.SUMMARY_SHEET]
        for row in range(4, 18):
            for col in "DEF":
                assert ws[f"{col}{row}"].value.startswith("=IFERROR("), f"{col}{row} unguarded"
        # PARTNERSHIPS is row 7; N/A means weight 0 and the rest scale back up.
        assert ws["C7"].value == 0.0
        assert abs(sum(ws[f"C{r}"].value for r in range(4, 18)) - 1.0) < 1e-9

    def test_score_column_is_a_formula_not_a_number(self):
        """The workbook must recalculate when a rating is changed in Excel."""
        from app import pitch_scorecard_xlsx as px
        ws = self._built()[px.CHECKLIST_SHEET]
        for row in range(4, 61):
            value = ws[f"I{row}"].value
            assert isinstance(value, str) and value.startswith("=IF(H")

    def test_unrated_criterion_exports_a_blank_rating(self):
        from openpyxl import load_workbook
        from app import pitch_scorecard_xlsx as px
        summary = _sample_summary()
        summary["categories"][0]["criteria"][0]["rating"] = None
        ws = load_workbook(px.build_workbook(summary), data_only=False)[px.CHECKLIST_SHEET]
        assert ws["H4"].value in ("", None)  # openpyxl reads "" back as None
        # The formula stays, so Excel shows "—" until somebody picks a rating.
        assert ws["I4"].value.startswith("=IF(H4")


# ── Model response validation ────────────────────────────────────────────────
class TestResponseParsing:
    """Whatever the model returns, nothing invented may reach the scorecard."""

    def _cat(self, no=1):
        return next(c for c in ps.load_rubric()["categories"] if c["no"] == no)

    def _parse(self, payload, counts=None):
        from app import pitch_scorecard_llm as llm
        import json as _json
        return llm._parse_response(
            _json.dumps(payload), self._cat(),
            counts or {"Deck.pptx": 40}, "Deck.pptx")

    def _entry(self, cid, **over):
        base = {"id": cid, "rating": "Best Practice", "evidence_status": "found",
                "slides": [{"doc": "Deck.pptx", "slide": 3}], "quote": "q",
                "comment": "c", "to_improve": "", "confidence": "high"}
        base.update(over)
        return base

    def test_happy_path(self):
        answers, problems = self._parse(
            {"criteria": [self._entry(f"1.{i}") for i in range(1, 5)]})
        assert len(answers) == 4
        assert problems == []
        assert answers["1.1"]["score"] == 2
        assert answers["1.1"]["source"] == ps.SOURCE_AI

    def test_strips_markdown_fence(self):
        from app import pitch_scorecard_llm as llm
        raw = '```json\n{"criteria": [{"id": "1.1", "rating": "Does Not Meet"}]}\n```'
        answers, _ = llm._parse_response(raw, self._cat(), {"Deck.pptx": 40}, "Deck.pptx")
        assert answers["1.1"]["rating"] == ps.RATING_FAIL

    def test_drops_slide_beyond_the_deck(self):
        """A cited slide 99 in a 40-slide deck is invented; it must not survive."""
        answers, problems = self._parse({"criteria": [
            self._entry("1.1", slides=[{"doc": "Deck.pptx", "slide": 99},
                                       {"doc": "Deck.pptx", "slide": 12}])]})
        assert answers["1.1"]["slides"] == [{"doc": "Deck.pptx", "slide": 12}]
        assert any("dropped citation" in p for p in problems)

    def test_rejects_invalid_rating(self):
        answers, problems = self._parse({"criteria": [
            self._entry("1.1", rating="Excellent")]})
        assert "1.1" not in answers
        assert any("unrecognised rating" in p for p in problems)

    def test_ignores_unknown_criterion_id(self):
        answers, problems = self._parse({"criteria": [
            self._entry("1.1"), self._entry("99.1")]})
        assert set(answers) == {"1.1"}
        assert any("unknown criterion" in p for p in problems)

    def test_reports_missing_criteria(self):
        _, problems = self._parse({"criteria": [self._entry("1.1")]})
        assert any("no rating returned" in p for p in problems)

    def test_resolves_mis_cased_doc_name(self):
        answers, _ = self._parse({"criteria": [
            self._entry("1.1", slides=[{"doc": "deck.PPTX", "slide": 3}])]})
        assert answers["1.1"]["slides"] == [{"doc": "Deck.pptx", "slide": 3}]

    def test_bare_int_citation_gets_the_default_doc(self):
        answers, _ = self._parse({"criteria": [self._entry("1.1", slides=[5])]})
        assert answers["1.1"]["slides"] == [{"doc": "Deck.pptx", "slide": 5}]

    def test_junk_response_is_survivable(self):
        from app import pitch_scorecard_llm as llm
        answers, problems = llm._parse_response(
            "I'm sorry, I can't help with that.", self._cat(),
            {"Deck.pptx": 40}, "Deck.pptx")
        assert answers == {}
        assert any("no JSON" in p for p in problems)

    def test_unknown_confidence_defaults_to_medium(self):
        answers, _ = self._parse({"criteria": [
            self._entry("1.1", confidence="extremely sure")]})
        assert answers["1.1"]["confidence"] == "medium"


class TestDigestParsing:
    def test_keeps_only_listed_slides(self):
        from app import pitch_scorecard_llm as llm
        raw = ("1 | Title slide | none\n"
               "2 | Exec summary | exec-summary, pov\n"
               "77 | Invented | pricing\n"
               "garbage line without a pipe\n")
        out = llm._parse_digest(raw, {1, 2, 3})
        assert set(out) == {1, 2}
        assert "exec-summary" in out[2]

    def test_tolerates_slide_prefix_and_bullets(self):
        from app import pitch_scorecard_llm as llm
        out = llm._parse_digest("- Slide 4 | Pricing table | pricing, placeholder-text",
                                {4})
        assert out[4].startswith("Pricing table")


# ── The deck is never sampled ────────────────────────────────────────────────
class TestWholeDeckCoverage:
    """Every criterion must be judged against every slide.

    The failure this guards against is subtle: most of the rubric is answered by
    absence ("no RACI anywhere"), so a model shown only part of the deck reports
    a confident, slide-cited gap that does not exist.
    """

    def _blocks(self, n_slides=40, chars=500):
        from app import pitch_scorecard_llm as llm
        slides = {"Deck.pptx": {i: "x" * chars for i in range(1, n_slides + 1)}}
        return llm.build_deck_blocks(slides, {})

    def test_every_slide_becomes_a_block(self):
        blocks = self._blocks(n_slides=115)
        assert len(blocks) == 115
        for i in range(1, 116):
            assert any(f"Slide {i} =" in label for label, _ in blocks), f"slide {i} missing"

    def test_slide_text_is_not_truncated(self):
        blocks = self._blocks(n_slides=3, chars=9000)
        assert all(len(text) == 9000 for _, text in blocks)

    def test_unpaged_documents_are_included(self):
        from app import pitch_scorecard_llm as llm
        blocks = llm.build_deck_blocks(
            {"Deck.pptx": {1: "slide one"}}, {"RFP.docx": "the rfp body text"})
        rendered = "\n".join(f"{a}{b}" for a, b in blocks)
        assert "the rfp body text" in rendered
        assert "no slide numbers" in rendered

    def test_deck_that_fits_is_one_part(self):
        from app import pitch_scorecard_llm as llm
        parts = llm.split_into_parts(self._blocks(n_slides=100, chars=1000))
        assert len(parts) == 1, "a deck within budget must not be split"

    def test_oversized_deck_splits_without_losing_a_slide(self):
        from app import pitch_scorecard_llm as llm
        blocks = self._blocks(n_slides=60, chars=1000)
        parts = llm.split_into_parts(blocks, char_budget=10_000)
        assert len(parts) > 1
        flattened = [b for part in parts for b in part]
        assert flattened == blocks, "splitting must preserve every block, in order"

    def test_split_parts_respect_the_budget(self):
        from app import pitch_scorecard_llm as llm
        blocks = self._blocks(n_slides=60, chars=1000)
        parts = llm.split_into_parts(blocks, char_budget=10_000)
        for part in parts:
            size = sum(len(a) + len(b) + 2 for a, b in part)
            # One oversized block can exceed the target on its own; the budget
            # bounds the packing, not a single indivisible slide.
            assert size <= 10_000 * 1.5 or len(part) == 1


class TestMergeVerdicts:
    """Merging a criterion's verdicts across parts of a split deck."""

    def _v(self, rating, slides=()):
        return {"rating": rating, "score": ps.RATING_SCORES[rating],
                "slides": [{"doc": "D.pptx", "slide": s} for s in slides]}

    def test_first_verdict_is_kept_when_alone(self):
        from app import pitch_scorecard_llm as llm
        v = self._v(ps.RATING_FAIL)
        assert llm.merge_verdicts(None, v) == v

    def test_evidence_beats_absence(self):
        """A 'not found' in part 1 must never override a 'found' in part 3."""
        from app import pitch_scorecard_llm as llm
        merged = llm.merge_verdicts(self._v(ps.RATING_FAIL, [2]),
                                    self._v(ps.RATING_BEST, [47]))
        assert merged["rating"] == ps.RATING_BEST

    def test_order_does_not_matter(self):
        from app import pitch_scorecard_llm as llm
        a = llm.merge_verdicts(self._v(ps.RATING_BEST, [47]), self._v(ps.RATING_FAIL, [2]))
        b = llm.merge_verdicts(self._v(ps.RATING_FAIL, [2]), self._v(ps.RATING_BEST, [47]))
        assert a["rating"] == b["rating"] == ps.RATING_BEST

    def test_citations_from_both_parts_survive(self):
        from app import pitch_scorecard_llm as llm
        merged = llm.merge_verdicts(self._v(ps.RATING_MIN, [3]),
                                    self._v(ps.RATING_BEST, [47, 48]))
        slides = sorted(s["slide"] for s in merged["slides"])
        assert slides == [3, 47, 48]

    def test_not_applicable_yields_to_a_real_rating(self):
        from app import pitch_scorecard_llm as llm
        merged = llm.merge_verdicts(self._v(ps.RATING_NA), self._v(ps.RATING_FAIL, [9]))
        assert merged["rating"] == ps.RATING_FAIL


class TestOutputBudget:
    def test_scales_with_criterion_count(self):
        """A six-criterion category must not be cut off mid-JSON."""
        from app import pitch_scorecard_llm as llm
        rubric = ps.load_rubric()
        branding = next(c for c in rubric["categories"] if c["no"] == 6)
        structure = next(c for c in rubric["categories"] if c["no"] == 1)
        assert len(branding["criteria"]) == 6 and len(structure["criteria"]) == 4
        assert llm._output_tokens_for(branding) > llm._output_tokens_for(structure)

    def test_never_below_the_floor(self):
        from app import pitch_scorecard_llm as llm
        tiny = {"criteria": [{"id": "1.1"}]}
        assert llm._output_tokens_for(tiny) >= llm._OUTPUT_FLOOR


def _run_standalone():
    """pytest isn't in requirements.txt yet — allow a plain-python run."""
    classes = [TestRubric, TestSampleScorecard, TestBands, TestNormalizeRating,
               TestNotApplicable, TestWeightOverrides, TestWorkbookExport,
               TestResponseParsing, TestDigestParsing, TestWholeDeckCoverage,
               TestMergeVerdicts, TestOutputBudget]
    passed = failed = 0
    for cls in classes:
        instance = cls()
        for name in sorted(n for n in dir(cls) if n.startswith("test_")):
            try:
                getattr(instance, name)()
                passed += 1
                print(f"  PASS  {cls.__name__}.{name}")
            except Exception as exc:
                failed += 1
                detail = str(exc).encode("ascii", "backslashreplace").decode("ascii")
                print(f"  FAIL  {cls.__name__}.{name}: {detail}")
    print(f"\n{passed} passed, {failed} failed")
    return 1 if failed else 0


if __name__ == "__main__":
    sys.exit(_run_standalone())
