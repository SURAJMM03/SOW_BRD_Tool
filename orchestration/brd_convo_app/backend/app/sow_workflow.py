"""
sow_workflow.py — Assemble the section-by-section SOW engine's approved
content into a final .docx.

Two paths, mirroring brd_workflow.py's own template-fill vs from-scratch
split — kept as separate code here (not shared with brd_workflow.py) so
nothing in this new SOW path can regress the working BRD export:

  save_sow_docx_from_template()  — a real .docx template (the org default or
                                    a project's own custom upload) is copied
                                    and each section's content is inserted
                                    directly under the template's own heading
                                    for that section title (heading-anchored
                                    fill, same strategy as brd_workflow.py's
                                    save_docx_from_template).
  save_sow_docx_fallback()       — no real template file exists yet (only the
                                    built-in FALLBACK_SECTIONS structure) —
                                    builds a plain .docx from scratch via
                                    sow_export.render_sow_docx.

Both paths finish with a post-processing pass: Markdown table syntax is
converted to real Word tables styled to match the Bristlecone template's own
tables (navy header row, Table Grid borders — a SOW-specific styler, not
screen3_routes.py's BRD-styled one), and <IMAGE id="..."/> tags are resolved
into embedded pictures by reusing screen3_routes.py's generic, section-
agnostic image resolver as-is.
"""
from __future__ import annotations

import copy
import logging
import re
import shutil
from pathlib import Path
from typing import Dict, List, Optional

from app.sow_export import EXPORT_DIR, sanitize_client_filename, render_sow_docx

logger = logging.getLogger("SOWWorkflow")


def _replace_in_paragraphs(paragraphs, replacements: List[tuple]) -> None:
    """Substitute each (placeholder, value) pair inside a list of python-docx
    paragraphs, span-scoped to exactly the matched characters in whichever
    run(s) they fall in — never touching a run's text outside the match, and
    never touching run.font/rPr at all.

    This is deliberately NOT the "overwrite run.runs[0], blank the rest"
    pattern used by brd_workflow.py's own header-building code (the suspected
    source of that path's known em-dash corruption bug) — Word can split a
    placeholder string across arbitrary run boundaries, so a naive whole-run
    overwrite can eat characters that were never part of the placeholder."""
    for p in paragraphs:
        for placeholder, value in replacements:
            while True:
                runs = p.runs
                full_text = "".join(r.text for r in runs)
                idx = full_text.find(placeholder)
                if idx < 0:
                    break
                match_start, match_end = idx, idx + len(placeholder)

                offsets = []
                pos = 0
                for r in runs:
                    offsets.append((pos, pos + len(r.text)))
                    pos += len(r.text)

                first = True
                for r, (r_start, r_end) in zip(runs, offsets):
                    overlap_start = max(r_start, match_start)
                    overlap_end = min(r_end, match_end)
                    if overlap_start >= overlap_end:
                        continue
                    local_start = overlap_start - r_start
                    local_end = overlap_end - r_start
                    before = r.text[:local_start]
                    after = r.text[local_end:]
                    if first:
                        r.text = before + value + after
                        first = False
                    else:
                        r.text = before + after


def _substitute_header_placeholders(doc, client_name: str, project_name: str) -> None:
    """The imported SOW templates ship a literal, never-substituted
    "[Client Name] | [Project Name]" string in their header (and this walks
    the footer too, in case a custom-uploaded template repeats it there).
    python-docx's doc.paragraphs excludes header/footer content entirely, so
    this must be done as an explicit separate pass."""
    replacements = [
        ("[Client Name]", client_name or "<CONFIRM: Client Name>"),
        ("[Project Name]", project_name or "<CONFIRM: Project Name>"),
    ]
    for section in doc.sections:
        for container in (section.header, section.footer):
            _replace_in_paragraphs(container.paragraphs, replacements)
            for table in container.tables:
                for row in table.rows:
                    for cell in row.cells:
                        _replace_in_paragraphs(cell.paragraphs, replacements)


def _substitute_cover_subtitle(doc, project_name: str) -> None:
    """The cover page's subtitle paragraph — sitting between the Title-
    styled paragraph and the very first heading (e.g. the real Bristlecone
    template ships the literal text "Master Template" there) — is never
    touched by any other pass: _substitute_header_placeholders only walks
    header/footer, and _insert_section_content only fills content AFTER a
    heading, but this paragraph sits BEFORE the first one. Deliberately
    content-agnostic (doesn't match on "Master Template" specifically) so
    it also works on a client's own uploaded custom template, whatever its
    own placeholder subtitle text happens to be."""
    paras = list(doc.paragraphs)
    title_idx = next(
        (i for i, p in enumerate(paras) if p.style and p.style.name == "Title"), None
    )
    if title_idx is None:
        return
    for p in paras[title_idx + 1:]:
        style = p.style.name if p.style else ""
        if style.startswith("Heading"):
            return  # reached the first heading — no subtitle paragraph exists
        if p.text.strip():
            new_text = project_name.strip() if project_name and project_name.strip() else "<CONFIRM: Project Name>"
            if not p.runs:
                p.add_run()
            p.runs[0].text = new_text
            for extra in p.runs[1:]:
                extra.text = ""
            return


def _find_heading_idx(all_paras, title: str) -> int:
    """Find a Heading-/Title-styled paragraph whose text matches `title`
    (exact first, then partial), mirroring brd_workflow.py's own matcher."""
    text_lower = (title or "").strip().lower()
    if not text_lower:
        return -1
    for i, p in enumerate(all_paras):
        style = p.style.name if p.style else ""
        if style.startswith("Heading") or style == "Title":
            if p.text.strip().lower() == text_lower:
                return i
    for i, p in enumerate(all_paras):
        style = p.style.name if p.style else ""
        if style.startswith("Heading"):
            if text_lower and text_lower in p.text.strip().lower():
                return i
    return -1


def _capture_body_rpr(doc, body_children: List, anchor_pos: int):
    """Read-only scan (before anything is removed) of the placeholder content
    directly under a heading, to find one representative run's <w:rPr> (font
    name/size/color/bold) so freshly-inserted content can be made to look
    like the template's own body text instead of an unstyled default run.
    Stops at the next heading, same boundary the removal loop respects."""
    import docx
    from docx.oxml.ns import qn

    for elem in body_children[anchor_pos + 1:]:
        tag = elem.tag.split("}")[-1]
        if tag != "p":
            continue
        p_obj = docx.text.paragraph.Paragraph(elem, doc)
        style = p_obj.style.name if p_obj.style else ""
        if style.startswith("Heading") or style == "Title":
            break
        for run in p_obj.runs:
            rpr = run._element.find(qn("w:rPr"))
            if rpr is not None:
                return copy.deepcopy(rpr)
    return None


def _clear_placeholder_body(doc, all_paras, heading_idx: int):
    """Remove placeholder paragraphs/tables between the given heading and the
    very next heading, at ANY level (see the caller's own note on why "any
    level", not "same-or-higher" — the real Bristlecone template's
    parent/child headings have no body of their own between them). Returns
    (anchor_element, captured_rpr) so callers can insert their own content
    right after the heading, styled like whatever placeholder body used to
    be there.

    Unlike brd_workflow.py's BRD template (whose in-place tables are static,
    reusable content like a Revision History table), every table under a
    SOW template section is a blank *example* table (e.g. "4.6 Integration
    Scope"'s two placeholder rows) that must be replaced, not left
    duplicated alongside new content."""
    anchor_elem = all_paras[heading_idx]._element
    body = doc.element.body
    body_children = list(body)
    anchor_pos = body_children.index(anchor_elem)

    # Capture the placeholder body's own run formatting BEFORE it's removed,
    # so whatever replaces it can match the template's own font/size/color
    # rather than falling back to an unstyled default run.
    captured_rpr = _capture_body_rpr(doc, body_children, anchor_pos)

    j = anchor_pos + 1
    while j < len(body_children):
        elem = body_children[j]
        tag = elem.tag.split("}")[-1]
        if tag == "p":
            import docx
            p_obj = docx.text.paragraph.Paragraph(elem, doc)
            style = p_obj.style.name if p_obj.style else ""
            if style.startswith("Heading") or style == "Title":
                break
            body.remove(elem)
            body_children = list(body)
            continue
        if tag == "tbl":
            body.remove(elem)
            body_children = list(body)
            continue
        j += 1

    return anchor_elem, captured_rpr


def _append_markdown_runs(p_elem, text: str, captured_rpr) -> None:
    """Append one or more <w:r> runs for `text` to the OOXML paragraph
    element `p_elem`, honoring **bold** spans (and dropping stray
    backticks) via sow_markdown.split_bold_segments — instead of writing
    the whole line as a single plain run, which is how `**bold**` markdown
    (the generation prompt explicitly tells the model to use it for
    emphasis) used to leak into the .docx as literal asterisks."""
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    from app.sow_markdown import split_bold_segments

    for seg_text, is_bold in (split_bold_segments(text) or [("", False)]):
        r = OxmlElement("w:r")
        rpr = copy.deepcopy(captured_rpr) if captured_rpr is not None else None
        if is_bold:
            if rpr is None:
                rpr = OxmlElement("w:rPr")
            rpr.append(OxmlElement("w:b"))
        if rpr is not None:
            r.append(rpr)
        t = OxmlElement("w:t")
        t.text = seg_text
        t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
        r.append(t)
        p_elem.append(r)


_IN_PROGRESS_TEXT = "⏳ In Progress — this section has not yet been drafted/approved."
_IN_PROGRESS_AMBER = "854F0B"      # matches sow-workflow.html's own --amber token
_IN_PROGRESS_SHADE = "FFF4E5"      # matches sow-workflow.html's own --amber-soft token


def _insert_in_progress_notice(doc, all_paras, heading_idx: int) -> None:
    """Replace a section's placeholder body with a short, visually-flagged
    notice (bold amber text on a light amber highlight, same brand colors
    the app's own sidebar/badges already use) instead of leaving the
    template's stale example content sitting under a heading that was never
    actually drafted — lets a partial download clearly show what's real
    content vs. not-yet-approved, rather than something a reader could
    mistake for a real (if generic) answer."""
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement

    anchor_elem, _captured_rpr = _clear_placeholder_body(doc, all_paras, heading_idx)

    p = OxmlElement("w:p")
    pPr = OxmlElement("w:pPr")
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), _IN_PROGRESS_SHADE)
    pPr.append(shd)
    p.append(pPr)

    r = OxmlElement("w:r")
    rPr = OxmlElement("w:rPr")
    rPr.append(OxmlElement("w:b"))
    color = OxmlElement("w:color")
    color.set(qn("w:val"), _IN_PROGRESS_AMBER)
    rPr.append(color)
    r.append(rPr)
    t = OxmlElement("w:t")
    t.text = _IN_PROGRESS_TEXT
    t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    r.append(t)
    p.append(r)

    anchor_elem.addnext(p)


def _insert_section_content(doc, all_paras, heading_idx: int, content: str) -> None:
    """Clear placeholder body paragraphs/tables after the heading (up to the
    very next heading, at any level) and insert the drafted content as plain
    paragraphs. Markdown tables and <IMAGE id="..."/> tags are left as literal
    text here — the caller's post-processing pass converts them afterward."""
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement

    anchor_elem, captured_rpr = _clear_placeholder_body(doc, all_paras, heading_idx)

    def _add_paragraph(text: str, list_style: bool) -> None:
        nonlocal insert_after
        p = OxmlElement("w:p")
        pPr = OxmlElement("w:pPr")
        if list_style:
            pStyle = OxmlElement("w:pStyle")
            pStyle.set(qn("w:val"), "ListParagraph")
            pPr.append(pStyle)
        p.append(pPr)
        _append_markdown_runs(p, text, captured_rpr)
        insert_after.addnext(p)
        insert_after = p
        if captured_rpr is None:
            # No representative run formatting was found under this heading
            # (a genuinely empty placeholder) — fall back to applying the
            # document's own Normal style font via the high-level API so the
            # run(s) still aren't left at Word's bare, unstyled default.
            import docx
            para_obj = docx.text.paragraph.Paragraph(p, doc)
            try:
                normal_font = doc.styles["Normal"].font
                for run_obj in para_obj.runs:
                    if normal_font.name:
                        run_obj.font.name = normal_font.name
                    if normal_font.size:
                        run_obj.font.size = normal_font.size
                    if normal_font.color and normal_font.color.rgb:
                        run_obj.font.color.rgb = normal_font.color.rgb
                    if normal_font.bold is not None and run_obj.font.bold is None:
                        run_obj.font.bold = normal_font.bold
            except Exception:
                pass

    insert_after = anchor_elem
    lines = [ln.strip() for ln in (content or "").split("\n")]
    i = 0
    while i < len(lines):
        line = lines[i]
        if not line:
            i += 1
            continue
        # A Markdown table block must land in ONE paragraph (its lines joined
        # by literal "\n") — the post-export table converter
        # (_convert_markdown_tables_bristlecone) detects a table by matching
        # the WHOLE block against a single paragraph's text; splitting each
        # row into its own paragraph (as this loop used to do for every line)
        # meant no single paragraph ever contained the full block, so the
        # table silently never converted and was left as raw "| a | b |" text.
        if line.startswith("|") and line.count("|") >= 2:
            table_lines = []
            while i < len(lines) and lines[i].startswith("|") and lines[i].count("|") >= 2:
                table_lines.append(lines[i])
                i += 1
            _add_paragraph("\n".join(table_lines), list_style=False)
            continue
        if line.startswith(("- ", "* ", "• ")):
            _add_paragraph(line[2:].strip(), list_style=True)
        else:
            _add_paragraph(line, list_style=False)
        i += 1


def save_sow_docx_from_template(
    sections: List[Dict],
    template_path: Path,
    client_name: str,
    project_name: str = "",
) -> Path:
    """Fill a real SOW .docx template with approved section content, anchored
    by matching each section's title against the template's own headings."""
    from docx import Document

    EXPORT_DIR.mkdir(parents=True, exist_ok=True)
    safe_client = sanitize_client_filename(client_name) or "Client"
    out_path = EXPORT_DIR / f"{safe_client}_SOW.docx"
    shutil.copy2(str(template_path), str(out_path))

    doc = Document(str(out_path))
    unmatched: List[Dict] = []

    _substitute_header_placeholders(doc, client_name, project_name)
    _substitute_cover_subtitle(doc, project_name)
    # Must run before the section-fill loop below replaces/removes the
    # template's own example tables (e.g. "4.6 Integration Scope"'s blank
    # placeholder rows) — this is the last point at which their original
    # run-level font size can still be read.
    table_body_pt = _capture_table_body_font_pt(doc)

    for sec in sections:
        content = (sec.get("content") or "").strip()
        all_paras = list(doc.paragraphs)
        h_idx = _find_heading_idx(all_paras, sec["title"])
        if h_idx < 0:
            if content:
                unmatched.append(sec)
            continue
        if content:
            _insert_section_content(doc, all_paras, h_idx, content)
        elif sec.get("has_own_content", True):
            # No approved content yet for a heading that's actually supposed
            # to have its own body — a partial download should say so
            # plainly rather than leaving the template's own stale example
            # placeholder sitting there looking like content.
            _insert_in_progress_notice(doc, all_paras, h_idx)
        # else: a container heading with no body slot of its own (its real
        # content lives in child sub-sections, which are handled by their
        # own entries in `sections`) — nothing to insert, nothing to flag.

    # Any section whose title didn't match a heading in the template (should
    # be rare, since the section list itself came from this template's TOC)
    # is appended at the end so content is never silently dropped.
    if unmatched:
        doc.add_page_break()
        for sec in unmatched:
            doc.add_heading(sec["title"], level=min(max(sec.get("level", 1), 1), 4))
            for line in (sec.get("content") or "").split("\n"):
                if line.strip():
                    doc.add_paragraph(line.strip())
        logger.info("SOW export: %d section(s) had no matching template heading — appended at end",
                    len(unmatched))

    doc.save(str(out_path))

    _post_process(out_path, table_body_pt=table_body_pt)
    return out_path


def save_sow_docx_fallback(sections: List[Dict], client_name: str) -> Path:
    """No real template file exists yet — build a plain .docx from scratch."""
    parts = []
    for sec in sections:
        content = (sec.get("content") or "").strip()
        if not content and not sec.get("has_own_content", True):
            continue  # container heading with no body of its own — nothing to show
        level = min(max(sec.get("level", 1), 1), 4)
        body = content or _IN_PROGRESS_TEXT
        parts.append(f"{'#' * level} {sec['title']}\n\n{body}")
    body_text = "\n\n".join(parts)

    out_path = render_sow_docx(body_text, "Statement of Work", client_name=client_name)
    _post_process(out_path)
    return out_path


# Bristlecone brand navy — matches build_sow_template.py's own table header
# shading, so a generated table looks identical to the template's own example
# tables (Table Grid style, navy header row, white bold header text) rather
# than screen3_routes.py's BRD-styled "Light Grid Accent 1".
_BRAND_NAVY = "1B3A6B"

# The separator row's dash-run must tolerate an alignment colon on either
# side (`:---`, `---:`, `:---:`) — without this, any LLM-generated table
# using Markdown column-alignment syntax (common, and confirmed in a real
# export's 13.2 Rate Card table) failed this block-detection regex entirely
# and the whole table was left as literal "| a | b |" text, never reaching
# the per-paragraph conversion loop below at all.
_TABLE_BLOCK_RE = re.compile(r'\|(?:[^\n\|]*\|)+\n\|(?:\s*:?-+:?\s*\|)+\n(?:\|(?:[^\n\|]*\|)+\n)*')
_SEP_ROW_RE = re.compile(r'^\|[\s\-:|]+\|$')


def _capture_table_body_font_pt(doc) -> Optional[float]:
    """Read the template's own example tables (still intact at this point,
    before the section-fill loop replaces/removes them) for a body-row run
    with an explicit font size, so generated tables can match it instead of a
    hardcoded value. Returns None if no example table has one — callers fall
    back to the prior hardcoded default (10pt), so this is a pure enhancement
    with no regression when a template's tables rely on style-level sizing."""
    for table in doc.tables:
        for row in table.rows[1:]:
            for cell in row.cells:
                for p in cell.paragraphs:
                    for run in p.runs:
                        if run.font.size:
                            return run.font.size.pt
    return None


def _set_cell_navy_header(cell) -> None:
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement

    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), _BRAND_NAVY)
    tcPr.append(shd)
    for p in cell.paragraphs:
        for run in p.runs:
            run.bold = True
            run.font.color.rgb = None
            rPr = run._element.get_or_add_rPr()
            color = OxmlElement("w:color")
            color.set(qn("w:val"), "FFFFFF")
            rPr.append(color)


def _set_cell_markdown_text(cell, text: str, font_pt: Optional[float] = None) -> None:
    """Write `text` into `cell` as one or more runs, honoring **bold** spans
    (and dropping stray backticks) via sow_markdown.split_bold_segments —
    python-docx's plain `cell.text = text` assignment always creates exactly
    one run and can't express bold at all, which is how markdown emphasis
    leaked into table cells (e.g. a real export's literal
    "`<CONFIRM: contract ID>`") verbatim."""
    from docx.shared import Pt
    from app.sow_markdown import split_bold_segments

    cell.text = ""  # resets to a single empty paragraph/run
    para = cell.paragraphs[0]
    for r in list(para.runs):
        r._element.getparent().remove(r._element)
    for seg_text, is_bold in (split_bold_segments(text) or [("", False)]):
        run = para.add_run(seg_text)
        if is_bold:
            run.bold = True
        if font_pt is not None:
            run.font.size = Pt(font_pt)


def _convert_markdown_tables_bristlecone(docx_path: Path, table_body_pt: Optional[float] = None) -> None:
    """Convert Markdown table syntax in the saved docx into real Word tables
    styled to match the Bristlecone template's own tables (navy header row,
    white bold header text, Table Grid borders) — a SOW-specific styling pass
    so exported tables look identical to the branded template's example
    tables, rather than reusing screen3_routes.py's BRD-styled converter."""
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    try:
        doc = Document(str(docx_path))
        modified = False
        paragraphs_to_process = []

        for para in doc.paragraphs:
            text = para.text
            if "|" in text and _TABLE_BLOCK_RE.search(text):
                paragraphs_to_process.append((para, text))

        for para, text in reversed(paragraphs_to_process):
            parts = re.split(f"({_TABLE_BLOCK_RE.pattern})", text)
            p_element = para._element
            p_parent = p_element.getparent()
            insert_position = list(p_parent).index(p_element)

            for part in parts:
                if not _TABLE_BLOCK_RE.match(part):
                    continue
                # A block may contain multiple adjacent tables — split into
                # (header, rows) groups so they don't merge into one table.
                blk = [ln for ln in part.strip().split("\n") if ln.strip()]
                groups = []
                gi = 0
                while gi < len(blk):
                    if gi + 1 < len(blk) and _SEP_ROW_RE.match(blk[gi + 1].strip()):
                        headers = [c.strip() for c in blk[gi].split("|") if c.strip()]
                        rows = []
                        k = gi + 2
                        while k < len(blk):
                            if _SEP_ROW_RE.match(blk[k].strip()):
                                break
                            if k + 1 < len(blk) and _SEP_ROW_RE.match(blk[k + 1].strip()):
                                break
                            cells = [c.strip() for c in blk[k].split("|") if c.strip()]
                            if len(cells) == len(headers):
                                rows.append(cells)
                            k += 1
                        if headers:
                            groups.append((headers, rows))
                        gi = k
                    else:
                        gi += 1

                for headers, rows in groups:
                    if not headers:
                        continue
                    tbl = doc.add_table(rows=len(rows) + 1, cols=len(headers))
                    tbl.style = "Table Grid"
                    hdr_cells = tbl.rows[0].cells
                    for col_idx, hdr_text in enumerate(headers):
                        if col_idx < len(hdr_cells):
                            _set_cell_markdown_text(hdr_cells[col_idx], hdr_text)
                            for p in hdr_cells[col_idx].paragraphs:
                                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                            _set_cell_navy_header(hdr_cells[col_idx])
                    for row_idx, row_data in enumerate(rows):
                        row_cells = tbl.rows[row_idx + 1].cells
                        for col_idx, cell_text in enumerate(row_data):
                            if col_idx < len(row_cells):
                                _set_cell_markdown_text(row_cells[col_idx], cell_text, font_pt=table_body_pt or 10)
                                for p in row_cells[col_idx].paragraphs:
                                    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                    tbl_element = tbl._element
                    tbl_element.getparent().remove(tbl_element)
                    p_parent.insert(insert_position, tbl_element)
                    insert_position += 1
                    modified = True

            p_parent.remove(p_element)

        if modified:
            doc.save(str(docx_path))
            logger.info("Converted Markdown tables to Bristlecone-styled Word tables: %s", docx_path.name)
    except Exception as exc:
        logger.warning("Failed to post-process tables in %s: %s", docx_path.name, exc)


def _rebuild_toc(docx_path: Path) -> None:
    """Insert a real, auto-updating Word TOC field under the document's
    "Table of Contents" heading (if present) so section-by-section navigation
    works properly, matching how Word's own References -> Table of Contents
    would populate it. brd_workflow.py's helpers are fully generic (they just
    search for a heading whose text contains "table of contents" — no BRD-
    specific coupling), so they're reused as-is rather than duplicated."""
    from docx import Document

    try:
        from app.brd_workflow import _rebuild_table_of_contents, _mark_fields_dirty
    except ImportError:
        from brd_workflow import _rebuild_table_of_contents, _mark_fields_dirty

    try:
        doc = Document(str(docx_path))
        _rebuild_table_of_contents(doc)
        _mark_fields_dirty(doc)
        doc.save(str(docx_path))
    except Exception as exc:
        logger.warning("Could not rebuild Table of Contents for %s: %s", docx_path.name, exc)


def _post_process(docx_path: Path, table_body_pt: Optional[float] = None) -> None:
    """Convert Markdown tables to Bristlecone-branded Word tables, resolve
    <IMAGE id=.../> tags into embedded pictures (that resolver is generic/
    section-agnostic, so screen3_routes.py's own implementation is reused
    as-is rather than duplicated), then rebuild the Table of Contents field
    for section-by-section navigation."""
    _convert_markdown_tables_bristlecone(docx_path, table_body_pt=table_body_pt)

    try:
        from app.screen3_routes import _resolve_images_in_docx
    except ImportError:
        from screen3_routes import _resolve_images_in_docx

    missing = _resolve_images_in_docx(docx_path)
    if missing:
        logger.warning("SOW export %s: %d image(s) could not be embedded: %s",
                        docx_path.name, len(missing), missing)

    _rebuild_toc(docx_path)


def build_sow_docx(project_id: str, sections: List[Dict], client_name: str,
                    template_docx_path: Optional[Path], project_name: str = "") -> Path:
    """Entry point used by the export endpoint: pick the template-fill path
    if a real template file resolved for this project, else the fallback
    from-scratch builder."""
    if template_docx_path and template_docx_path.exists():
        return save_sow_docx_from_template(sections, template_docx_path, client_name, project_name)
    return save_sow_docx_fallback(sections, client_name)
